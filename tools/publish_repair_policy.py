from __future__ import annotations

import hashlib
import io
import json
from urllib.parse import quote, urlparse

from docx import Document
from docx.shared import Cm, Pt, RGBColor
import httpx

from content_studio.auth import AuthorSession
from fleet.repair_policy import POLICY_PATH, PUBLICATION_PATH, repair_policy
from tools.cloud import Cloud, CONFIG, GRAPH, save_state

DOCX_NAME = "Caldova-Repair-Policy.docx"


def build_document(policy: dict) -> bytes:
    document = Document()
    document.core_properties.title = policy["title"]
    document.core_properties.subject = "New genuine OEM parts, quotation compliance and repair approval"
    document.core_properties.author = policy["owner"]
    document.core_properties.version = policy["version"]
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    style = document.styles["Normal"]
    style.font.name = "Aptos"
    style.font.size = Pt(10.5)
    document.add_heading("CALDOVA DRIVE", 0)
    document.add_heading("Repair Policy", 1)
    document.add_paragraph(f"{policy['id']} | Version {policy['version']} | Effective {policy['effective_date']}")
    document.add_paragraph(f"Policy owner: {policy['owner']}")
    executive = document.add_paragraph()
    executive.add_run("New genuine OEM replacement parts only. ").bold = True
    executive.add_run("A faster or cheaper offer cannot override parts compliance.")
    for clause in policy["clauses"]:
        heading = document.add_heading(f"{clause['id']} - {clause['title']}", 2)
        for run in heading.runs:
            run.font.color.rgb = RGBColor.from_string("234932")
        document.add_paragraph(clause["text"])
    document.add_heading("Quotation review checklist", 2)
    for item in [
        "Is the repair centre approved, and is the quote valid?",
        "Does the quote explicitly state whether parts are replaced?",
        "Is every replacement declared new, genuine OEM and vehicle-manufacturer approved?",
        "Are part identity, manufacturer, scope, price including VAT, start date, completion date and warranty stated?",
        "Are any cheaper/faster offers excluded or held for clarification? Record the clause and source statement.",
        "Has the operator approved an eligible offer, and has the garage subsequently confirmed the same terms?",
    ]:
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Illustrative decision", 2)
    document.add_paragraph(
        "A quotation offering new aftermarket bumper parts may be cheaper and faster, but it fails RP-02 and is excluded under RP-05. "
        "An explicitly new genuine OEM offer remains eligible for comparison. If an offer simply says 'replacement bumper', request RP-03 clarification; do not assume it is OEM."
    )
    section.footer.paragraphs[0].text = f"Caldova Drive | {policy['id']} | v{policy['version']} | Controlled repair policy"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def main():
    Cloud()
    policy = repair_policy()
    payload = build_document(policy)
    file = POLICY_PATH.parent / DOCX_NAME
    author = AuthorSession(CONFIG["report_recipient"])
    identity = author.verify()
    if identity["id"] != CONFIG["admin_object_id"]:
        raise RuntimeError("Only the configured Caldova operator can publish the controlled policy.")

    def request(method, url, body=None, *, content=None, content_type="application/json"):
        if not url.startswith(GRAPH + "/"):
            raise ValueError("Policy publication is restricted to Microsoft Graph.")
        response = httpx.request(
            method, url, json=body, content=content, timeout=90,
            headers={"Authorization": f"Bearer {author.token()}", "Content-Type": content_type},
        )
        response.raise_for_status()
        return response.json() if response.content else None

    drive = request("GET", GRAPH + "/me/drive?$select=id,webUrl,driveType")
    if drive["driveType"] != "business" or urlparse(drive["webUrl"]).hostname != "caldova08667473-my.sharepoint.com":
        raise RuntimeError("Refusing to publish outside the Caldova operator's Microsoft 365 document library.")
    root = f"{GRAPH}/drives/{drive['id']}"
    parent = "root"
    for name in ("Caldova Drive", "Policies"):
        children = []
        page_url = f"{root}/{parent}/children?$select=id,name,folder"
        while page_url:
            page = request("GET", page_url)
            children.extend(page["value"])
            page_url = page.get("@odata.nextLink")
        existing = next((item for item in children if item["name"] == name), None)
        if existing and "folder" not in existing:
            raise RuntimeError(f"A file conflicts with the policy folder {name}.")
        folder = existing or request("POST", f"{root}/{parent}/children", {
            "name": name, "folder": {}, "@microsoft.graph.conflictBehavior": "fail",
        })
        parent = "items/" + folder["id"]
    uploaded = request(
        "PUT", f"{root}/{parent}:/{quote(DOCX_NAME)}:/content",
        content=payload, content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    # Verify the downloaded file is the exact published Word document, without logging its short-lived URL.
    item = request("GET", f"{root}/items/{uploaded['id']}")
    with httpx.Client(timeout=60, follow_redirects=True) as client:
        downloaded = client.get(item["@microsoft.graph.downloadUrl"])
        if not downloaded.is_success:
            raise RuntimeError(f"SharePoint document verification returned HTTP {downloaded.status_code}.")
    digest = hashlib.sha256(payload).hexdigest()
    if hashlib.sha256(downloaded.content).hexdigest() != digest:
        raise RuntimeError("The published repair policy differs from the generated Word document.")
    file.write_bytes(payload)
    publication = {
        "policy_id": policy["id"], "version": policy["version"],
        "source_sha256": hashlib.sha256(POLICY_PATH.read_bytes()).hexdigest(),
        "document_sha256": digest, "document_url": uploaded["webUrl"],
        "library": "Caldova operator OneDrive for Business",
        "drive_id": drive["id"], "drive_item_id": uploaded["id"], "etag": uploaded["eTag"],
    }
    PUBLICATION_PATH.write_text(json.dumps(publication, indent=2), encoding="utf-8")
    save_state({"repair_policy": publication})
    print(f"Word policy: {file}")
    print(f"Caldova OneDrive for Business: {publication['document_url']}")
    print("Published Word bytes verified; the operator owns the document in their Caldova Microsoft 365 library.")


if __name__ == "__main__":
    main()
