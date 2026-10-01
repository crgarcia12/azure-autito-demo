from __future__ import annotations

from datetime import datetime
import hashlib
import html
import io
import json
import threading
import time
from urllib.parse import quote
import uuid

from docx import Document
from docx.shared import Pt, RGBColor
import httpx

from content_studio.auth import AuthorSession, personas
from fleet.storage import StateStore
from tools.cloud import CONFIG, GRAPH, ROOT


def validate_plan(plan: dict, people: dict[str, dict]) -> None:
    uuid.UUID(plan["runId"])
    if plan["tenantId"] != CONFIG["tenant_id"]:
        raise ValueError("This content plan belongs to a different tenant.")
    expected = json.loads((ROOT / "content.config.json").read_text(encoding="utf-8"))
    counts = (
        len(plan["conversations"]),
        sum(len(chat["messages"]) for chat in plan["conversations"]),
        len(plan["meetings"]), len(plan["documents"]),
    )
    if counts != (expected["conversations"], expected["conversations"] * expected["messages_per_conversation"], expected["meetings"], expected["documents"]):
        raise ValueError(f"Content volume differs from the approved batch: {counts}.")
    ids = [item["id"] for group in ("conversations", "meetings", "documents") for item in plan[group]]
    if len(ids) != len(set(ids)):
        raise ValueError("Every generated artifact must have a unique delivery identifier.")
    allowed = {upn.casefold() for upn in people}
    authors = {item["upn"].casefold() for item in expected["authors"]}
    observer = expected["observer"].casefold()
    for chat in plan["conversations"]:
        members = {upn.casefold() for upn in chat["members"]}
        if not members.issubset(allowed) or observer not in members or chat["creator"].casefold() not in members:
            raise ValueError("Every conversation must include only approved Caldova participants, including the observer and creator.")
        if len(chat["messages"]) != expected["messages_per_conversation"]:
            raise ValueError("A conversation differs from its approved message count.")
        for message in chat["messages"]:
            if message["author"].casefold() not in authors & members or not message["text"].strip():
                raise ValueError("A message has an unapproved author or empty content.")
    for meeting in plan["meetings"]:
        participants = {upn.casefold() for upn in meeting["attendees"]}
        if meeting["organizer"].casefold() not in authors or not participants.issubset(allowed) or observer not in participants:
            raise ValueError("A meeting has an unapproved participant.")
        start, end = datetime.fromisoformat(meeting["start"]), datetime.fromisoformat(meeting["end"])
        if start.tzinfo is None or end.tzinfo is None or end <= start:
            raise ValueError("Meeting dates must be timezone-aware and have positive duration.")
    for document in plan["documents"]:
        readers = {upn.casefold() for upn in document["readers"]}
        if document["author"].casefold() not in authors or not readers.issubset(allowed) or observer not in readers:
            raise ValueError("A document has an unapproved author or recipient.")


class GraphError(RuntimeError):
    def __init__(self, status: int, details: str):
        self.status = status
        super().__init__(f"Microsoft Graph HTTP {status}: {details[:1200]}")


class AuthorGraph:
    def __init__(self, author: AuthorSession) -> None:
        self.author = author
        self.http = httpx.Client(timeout=90)

    def request(self, method: str, path: str, body=None, *, content=None, content_type=None):
        if not path.startswith(("/me/", "/chats")):
            raise ValueError("Content Studio can access only the signed-in author's drive/calendar and the specified chats.")
        headers = {"Authorization": f"Bearer {self.author.token()}"}
        if content_type:
            headers["Content-Type"] = content_type
        for attempt in range(6):
            response = self.http.request(
                method, GRAPH + path, json=body, content=content, headers=headers
            )
            if response.status_code != 429:
                break
            if attempt == 5:
                raise GraphError(response.status_code, response.text)
            time.sleep(min(int(response.headers.get("Retry-After", "10")), 120))
        if response.is_error:
            raise GraphError(response.status_code, response.text)
        if response.status_code == 207 and any("error" in item for item in response.json().get("value", [])):
            raise GraphError(207, "One or more document recipients did not receive access.")
        return response.json() if response.content else None


def word_document(document: dict, source: dict) -> bytes:
    word = Document()
    word.core_properties.title = document["title"]
    word.core_properties.author = "Caldova Drive"
    normal = word.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    word.add_heading(document["title"], 0)
    word.add_paragraph(f"Caldova Drive | Reporting context: {source['reportDate']} | {source['reportTimezone']}")
    for section in document["sections"]:
        heading = word.add_heading(section["heading"], 1)
        for run in heading.runs:
            run.font.color.rgb = RGBColor.from_string("234932")
        for paragraph in section["paragraphs"]:
            word.add_paragraph(paragraph)
    output = io.BytesIO()
    word.save(output)
    return output.getvalue()


class Publisher:
    def __init__(self, plan: dict) -> None:
        self.plan = plan
        self.people = personas()
        validate_plan(plan, self.people)
        self.store = StateStore(ROOT / ".local" / "content-delivery.sqlite3")
        self.clients: dict[str, AuthorGraph] = {}
        self.run = plan["runId"]
        self.lease_error: Exception | None = None

    def client(self, upn: str) -> AuthorGraph:
        key = upn.casefold()
        if key not in self.clients:
            self.clients[key] = AuthorGraph(AuthorSession(upn))
        return self.clients[key]

    def preflight(self) -> None:
        authors = sorted({message["author"] for chat in self.plan["conversations"] for message in chat["messages"]})
        for upn in authors:
            person = self.client(upn).author.verify()
            print(f"Authenticated author: {person['displayName']} ({person['userPrincipalName']})", flush=True)
        document_owners = {document["author"] for document in self.plan["documents"]}
        for owner in document_owners:
            self.client(owner).request("GET", "/me/drive?$select=id,driveType")
        for organizer in {meeting["organizer"] for meeting in self.plan["meetings"]}:
            calendar = self.client(organizer).request("GET", "/me/calendar?$select=allowedOnlineMeetingProviders")
            if "teamsForBusiness" not in calendar.get("allowedOnlineMeetingProviders", []):
                raise RuntimeError(f"Teams meetings are not enabled for {organizer}. Nothing has been published.")

    def deliver(self, key: str, call, *, retry_safe: bool = False) -> dict:
        if self.lease_error is not None:
            raise RuntimeError("The content delivery lease was lost; no further requests will be sent.") from self.lease_error
        ledger_key = f"{self.run}/{key}"
        prior = self.store.get(ledger_key)
        if prior and prior["status"] == "completed":
            return prior["response"]
        if prior and prior["status"] == "sending" and not retry_safe:
            raise RuntimeError(
                f"{key} has an uncertain prior delivery. Reconcile its Graph result before retrying; "
                "Content Studio will not silently duplicate chats or messages."
            )
        self.store.put(ledger_key, {"status": "sending"})
        try:
            response = call()
        except GraphError as error:
            self.store.put(ledger_key, {"status": "rejected", "error": str(error)})
            raise
        self.store.put(ledger_key, {"status": "completed", "response": response})
        return response

    def documents(self) -> list[dict]:
        results = []
        folder_name = json.loads((ROOT / "content.config.json").read_text(encoding="utf-8"))["drive_folder"]
        for document in self.plan["documents"]:
            client = self.client(document["author"])
            try:
                folder = client.request("GET", "/me/drive/root:/" + quote(folder_name))
            except GraphError as error:
                if error.status != 404:
                    raise
                folder = client.request("POST", "/me/drive/root/children", {
                    "name": folder_name, "folder": {}, "@microsoft.graph.conflictBehavior": "fail",
                })
            filename = f"{self.plan['source']['reportDate']} - {document['title']}.docx"
            payload = word_document(document, self.plan["source"])
            uploaded = self.deliver(
                document["id"],
                lambda: client.request(
                    "PUT", f"/me/drive/items/{folder['id']}:/{quote(filename)}:/content",
                    content=payload,
                    content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
                retry_safe=True,
            )
            recipients = [
                {"email": upn} for upn in document["readers"]
                if upn.casefold() != document["author"].casefold()
            ]
            self.deliver(
                document["id"] + "/sharing",
                lambda: client.request("POST", f"/me/drive/items/{uploaded['id']}/invite", {
                    "requireSignIn": True, "sendInvitation": False, "roles": ["read"],
                    "recipients": recipients,
                }),
                retry_safe=True,
            )
            results.append({"title": document["title"], "webUrl": uploaded["webUrl"], "id": uploaded["id"]})
            print(f"Shared document: {document['title']}", flush=True)
        return results

    def meetings(self, documents: list[dict]) -> list[dict]:
        results = []
        for meeting in self.plan["meetings"]:
            client = self.client(meeting["organizer"])
            agenda = "<h2>Agenda</h2><ul>" + "".join(f"<li>{html.escape(item)}</li>" for item in meeting["agenda"]) + "</ul>"
            agenda += "<h2>Shared working documents</h2><ul>" + "".join(
                f'<li><a href="{html.escape(item["webUrl"], quote=True)}">{html.escape(item["title"])}</a></li>' for item in documents
            ) + "</ul>"
            event = self.deliver(
                meeting["id"],
                lambda: client.request("POST", "/me/events", {
                    "subject": meeting["subject"], "body": {"contentType": "HTML", "content": agenda},
                    "start": {"dateTime": datetime.fromisoformat(meeting["start"]).replace(tzinfo=None).isoformat(), "timeZone": "GMT Standard Time"},
                    "end": {"dateTime": datetime.fromisoformat(meeting["end"]).replace(tzinfo=None).isoformat(), "timeZone": "GMT Standard Time"},
                    "attendees": [
                        {"emailAddress": {"address": upn}, "type": "required"}
                        for upn in meeting["attendees"] if upn.casefold() != meeting["organizer"].casefold()
                    ],
                    "isOnlineMeeting": True, "onlineMeetingProvider": "teamsForBusiness",
                    "allowNewTimeProposals": False,
                    "transactionId": str(uuid.uuid5(uuid.UUID(self.run), meeting["id"])),
                }),
                retry_safe=True,
            )
            join_url = (event.get("onlineMeeting") or {}).get("joinUrl")
            if not event.get("isOnlineMeeting") or not join_url:
                raise RuntimeError(f"Graph did not create a Teams join link for {meeting['subject']}.")
            results.append({"id": event["id"], "subject": meeting["subject"], "joinUrl": join_url})
            print(f"Created Teams meeting: {meeting['subject']}", flush=True)
        return results

    def conversations(self, documents: list[dict], meetings: list[dict]) -> list[dict]:
        results = []
        for index, chat in enumerate(self.plan["conversations"]):
            creator = self.client(chat["creator"])
            created = self.deliver(
                chat["id"],
                lambda: creator.request("POST", "/chats", {
                    "chatType": "group", "topic": chat["topic"],
                    "members": [{
                        "@odata.type": "#microsoft.graph.aadUserConversationMember", "roles": ["owner"],
                        "user@odata.bind": f"{GRAPH}/users('{self.people[upn.casefold()]['id']}')",
                    } for upn in chat["members"]],
                }),
            )
            messages = []
            for number, message in enumerate(chat["messages"]):
                body = html.escape(message["text"])
                if number == 0:
                    document = documents[index % len(documents)]
                    body += f'<br><br>Working document: <a href="{html.escape(document["webUrl"], quote=True)}">{html.escape(document["title"])}</a>'
                if number == len(chat["messages"]) - 1:
                    meeting = meetings[index % len(meetings)]
                    body += f'<br><br>Next discussion: <a href="{html.escape(meeting["joinUrl"], quote=True)}">{html.escape(meeting["subject"])}</a>'
                author = self.client(message["author"])
                sent = self.deliver(
                    f"{chat['id']}/message-{number + 1}",
                    lambda: author.request(
                        "POST", f"/chats/{quote(created['id'], safe='')}/messages",
                        {"body": {"contentType": "html", "content": body}},
                    ),
                )
                sender = (sent.get("from") or {}).get("user", {}).get("id")
                if sender != self.people[message["author"].casefold()]["id"]:
                    raise RuntimeError("Graph attributed a message to an unexpected author.")
                messages.append(sent["id"])
                time.sleep(1.1)
            results.append({"topic": chat["topic"], "chatId": created["id"], "messageIds": messages})
            print(f"Published conversation: {chat['topic']} ({len(messages)} messages)", flush=True)
        return results

    def publish(self) -> dict:
        self.preflight()
        digest = hashlib.sha256(json.dumps(self.plan, sort_keys=True).encode()).hexdigest()
        key = f"{self.run}/plan"
        prior = self.store.get(key)
        if prior and prior["sha256"] != digest:
            raise RuntimeError("This run already has a different approved plan. Do not mutate a partially published run.")
        self.store.put(key, {"sha256": digest})
        with self.store.lease(f"content-{self.run}") as lease:
            stopping = threading.Event()

            def heartbeat():
                while not stopping.wait(15):
                    try:
                        lease.renew()
                    except Exception as error:
                        self.lease_error = error
                        return

            thread = threading.Thread(target=heartbeat, daemon=True)
            thread.start()
            try:
                documents = self.documents()
                meetings = self.meetings(documents)
                conversations = self.conversations(documents, meetings)
            finally:
                stopping.set()
                thread.join(timeout=35)
        if self.lease_error is not None:
            raise RuntimeError("Delivery lease renewal failed. Review the persisted receipts before continuing.") from self.lease_error
        result = {"runId": self.run, "tenantId": CONFIG["tenant_id"], "documents": documents, "meetings": meetings, "conversations": conversations}
        (ROOT / ".local" / "content-published.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
        return result
